# Разовый перенос НПС из Excel на сайт: «НПС GATE H&A.xlsx» (полные листы) + «Анкария GATE H&A.xlsx» (только статус и место жительства).
# Пишет папку-пакет: npcs.json и art/*.webp (целиком до 2000px и превью 720px). Дальше пакет забирает server/import-npcs.js.
#   python tools/npc-xlsx-to-json.py "<НПС.xlsx>" "<Анкария.xlsx>" <папка-пакета>
import io, json, os, re, sys
import openpyxl
from PIL import Image

SRC_A, SRC_B, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
os.makedirs(os.path.join(OUT, 'art'), exist_ok=True)

GROUPS = {'Сайд-кики': 'sidekick', 'Личные Сайд-кики': 'personal', 'Компаньоны': 'companion', 'Важные НПС': 'important'}
# имена, как их правильно писать (ответ мастера); ключ — как встречается во втором файле
CANON = {'Арчебальд Шварцерр': 'Арчибальд Шварцерр', 'Арамэ': 'Арамэ Клабситис', 'Кайл Лотнер': 'Каил Лотнер', 'Ревдарий Е.В.И.У.С': 'Ревдарий Евиус'}
SPEC_LEVELS = ['Новичок', 'Умелец', 'Эксперт', 'Мастер', 'Легенда']
WEAK_LEVELS = {'Плохо': 'Плохо', 'Некомпетентность': 'Некомпетентный', 'Некомпетентный': 'Некомпетентный', 'Ужасно': 'Ужасно', 'Ужастно': 'Ужасно'}
SLOT_WORDS = ('Слот владения', 'Слот техники', 'Слот дрессировки', 'Слот заклинания')
seen_levels = {'spec': set(), 'weak': set()}

def txt(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return re.sub(r'[ \t]+', ' ', str(v)).strip()

def num(v):
    try: return float(v) if v not in (None, '') else None
    except (TypeError, ValueError): return None

def save_art(blob, key):
    im = Image.open(io.BytesIO(blob))
    im = im.convert('RGBA') if im.mode in ('P', 'LA', 'RGBA') else im.convert('RGB')
    out = {}
    for suffix, side, q in (('', 2000, 90), ('-t', 720, 86)):
        c = im.copy()
        c.thumbnail((side, side), Image.LANCZOS)
        name = f'{key}{suffix}.webp'
        c.save(os.path.join(OUT, 'art', name), 'WEBP', quality=q, method=6)
        out['file' if not suffix else 'thumb'] = name
    return out

def cell(ws, r, c): return ws.cell(row=r, column=c).value

def parse_block(ws, r0, group, images):
    end = r0 + 99
    npc = {'group': group, 'info': [], 'lists': [], 'specs': [], 'weak': [], 'prof': [], 'profSlots': 0, 'passives': [], 'actives': [],
           'combat': {}, 'stats': {}, 'saves': [], 'resist': [], 'attacks': [], 'spells': [], 'spellSlots': 0, 'spellsTitle': '', 'arts': []}
    # верх: основная информация (G/H) и две колонки-списка (J/K, M/N) с заголовками в строке r0
    head2, head3 = txt(cell(ws, r0, 10)), txt(cell(ws, r0, 13))
    col2, col3 = [], []
    for r in range(r0 + 1, r0 + 14):
        k, v = txt(cell(ws, r, 7)), txt(cell(ws, r, 8))
        if k == 'Имя': npc['name'] = v
        elif k: npc['info'].append({'k': k, 'v': v})
        a, b = txt(cell(ws, r, 10)), txt(cell(ws, r, 11))
        if a: col2.append((a, b))
        a, b = txt(cell(ws, r, 13)), txt(cell(ws, r, 14))
        if a: col3.append((a, b))
    if head2 == 'Специализации':
        for a, b in col2:
            seen_levels['spec'].add(b)
            npc['specs'].append({'name': a, 'level': b})
        for a, b in col3:
            seen_levels['weak'].add(b)
            npc['weak'].append({'name': a, 'level': WEAK_LEVELS.get(b, b)})
    else:  # компаньоны: снаряжение/инвентарь; важные: профессии/квесты
        if head2: npc['lists'].append({'title': head2, 'items': [a + (f' — {b}' if b else '') for a, b in col2]})
        if head3: npc['lists'].append({'title': head3, 'items': [a + (f' — {b}' if b else '') for a, b in col3]})
    # середина: владения (G/H), пассивные (J/K), третья колонка (M/N: занятия и хобби / любимые вещи)
    hdr = next((r for r in range(r0, end) if txt(cell(ws, r, 7)) == 'Владение'), None)
    if hdr:
        third = txt(cell(ws, hdr, 13)); items3 = []
        for r in range(hdr + 1, hdr + 16):
            k, v = txt(cell(ws, r, 7)), txt(cell(ws, r, 8))
            if k in SLOT_WORDS: npc['profSlots'] += 1
            elif k: npc['prof'].append({'name': k, 'level': v})
            a, b = txt(cell(ws, r, 10)), txt(cell(ws, r, 11))
            if a: npc['passives'].append({'name': a, 'type': b})
            a, b = txt(cell(ws, r, 13)), txt(cell(ws, r, 14))
            if a: items3.append(a + (f' — {b}' if b else ''))
        if third: npc['lists'].append({'title': third, 'items': items3})
    # характеристики: 6 строк — слева боевые числа (A/B), справа атрибуты (C/D)
    r = next((r for r in range(r0, end) if txt(cell(ws, r, 1)) == 'Характеристики'), None)
    if r:
        for i in range(1, 7):
            k, v = txt(cell(ws, r + i, 1)), cell(ws, r + i, 2)
            if k: npc['combat'][k] = txt(v)
            k, v = txt(cell(ws, r + i, 3)), cell(ws, r + i, 4)
            if k: npc['stats'][k] = txt(v)
    r = next((r for r in range(r0, end) if txt(cell(ws, r, 1)) == 'Спасброски'), None)
    if r:
        for i in range(1, 6):
            k, v = txt(cell(ws, r + i, 1)), txt(cell(ws, r + i, 2))
            if k in ('Атаки',): break
            if k: npc['saves'].append({'stat': k, 'value': v})
            k, v = txt(cell(ws, r + i, 4)), txt(cell(ws, r + i, 5))
            if k: npc['resist'].append({'kind': k, 'value': v})
    # атаки до блока заклинаний/техник/дрессировок
    ra = next((r for r in range(r0, end) if txt(cell(ws, r, 1)) == 'Атаки'), None)
    rs = next((r for r in range(r0, end) if txt(cell(ws, r, 1)) in ('Заклинания и Техники', 'Дрессировки', 'Заклинания')), None)
    if ra:
        for r in range(ra + 1, (rs or ra + 6)):
            name = txt(cell(ws, r, 1))
            if not name: continue
            npc['attacks'].append({'name': name, 'dtype': txt(cell(ws, r, 2)), 'kind': txt(cell(ws, r, 3)), 'dmg': txt(cell(ws, r, 4)), 'note': txt(cell(ws, r, 5))})
    if rs:
        npc['spellsTitle'] = txt(cell(ws, rs, 1))
        for r in range(rs + 1, end):
            v = txt(cell(ws, r, 1))
            if not v: continue
            if v in SLOT_WORDS: npc['spellSlots'] += 1
            else: npc['spells'].append(v)
    # активные навыки: объединённые блоки G..N после заголовка «Активные навыки»
    rh = next((r for r in range(r0, end) if txt(cell(ws, r, 7)) == 'Активные навыки'), None)
    if rh:
        for r in range(rh + 1, end + 1):
            name = txt(cell(ws, r, 7))
            if not name: continue
            npc['actives'].append({'name': name, 'type': txt(cell(ws, r, 8)), 'desc': txt(cell(ws, r, 9)), 'cooldown': txt(cell(ws, r, 13)), 'cost': txt(cell(ws, r, 14))})
    # арты: картинки, привязанные к строкам этого блока (слева направо)
    mine = sorted([im for im in images if r0 <= im.anchor._from.row + 1 <= end], key=lambda im: (im.anchor._from.row, im.anchor._from.col))
    npc['_images'] = mine
    return npc

npcs = []
wa = openpyxl.load_workbook(SRC_A, data_only=True)
for sheet, group in GROUPS.items():
    ws = wa[sheet]
    imgs = list(getattr(ws, '_images', []))
    starts = [r for r in range(1, ws.max_row + 1) if txt(cell(ws, r, 7)) == 'Основная информация']
    for r0 in starts:
        n = parse_block(ws, r0, group, imgs)
        if n.get('name'): npcs.append(n)

# Аспект (дракон) — свой лист другой раскладки
ws = wa['Аспекты']
for r0 in [r for r in range(1, ws.max_row + 1) if txt(cell(ws, r, 1)) == 'Имя']:
    n = {'group': 'aspect', 'info': [], 'lists': [], 'specs': [], 'weak': [], 'prof': [], 'profSlots': 0, 'passives': [], 'actives': [],
         'combat': {}, 'stats': {}, 'saves': [], 'resist': [], 'attacks': [], 'spells': [], 'spellSlots': 0, 'spellsTitle': '', 'arts': []}
    for r in range(r0, r0 + 12):
        k, v = txt(cell(ws, r, 1)), txt(cell(ws, r, 2))
        if k == 'Имя': n['name'] = v
        elif k: n['info'].append({'k': k, 'v': v})
        k, v = txt(cell(ws, r, 3)), txt(cell(ws, r, 4))
        if k in ('БМ', 'Инициатива', 'Скорость'): n['combat'][k] = v
        elif k: n['stats'][k] = v
        k, v = txt(cell(ws, r, 5)), txt(cell(ws, r, 6))
        if k in ('Хиты', 'Броня', 'Уклонение'): n['combat'][k] = v
        elif k == 'Спасбросок': n['saves'].append({'stat': v, 'value': ''})
        elif k: n['resist'].append({'kind': v, 'value': k})
    n['_images'] = sorted(getattr(ws, '_images', []), key=lambda im: (im.anchor._from.row, im.anchor._from.col))
    npcs.append(n)

# второй файл: статус и место жительства из «Списка НПС»
wb = openpyxl.load_workbook(SRC_B, data_only=True)
extra = {}
for row in wb['Список НПС'].iter_rows(min_row=2):
    nm = txt(row[1].value)
    if nm: extra[CANON.get(nm, nm)] = {'status': txt(row[5].value), 'home': txt(row[7].value)}
missing_extra = []
for n in npcs:
    e = extra.pop(n['name'], None)
    if e: n.update(e)
    elif n['group'] in ('sidekick', 'personal'): missing_extra.append(n['name'])

# картинки
for i, n in enumerate(npcs):
    for j, im in enumerate(n.pop('_images')):
        n['arts'].append(save_art(im._data(), f'npc{i:02d}-{j}'))

with open(os.path.join(OUT, 'npcs.json'), 'w', encoding='utf-8') as f:
    json.dump(npcs, f, ensure_ascii=False, indent=1)

print('НПС:', len(npcs), {g: sum(1 for n in npcs if n['group'] == g) for g in ['sidekick', 'personal', 'companion', 'important', 'aspect']})
print('артов:', sum(len(n['arts']) for n in npcs), '· без арта:', [n['name'] for n in npcs if not n['arts']])
print('уровни специализаций:', sorted(seen_levels['spec']), '· слабостей:', sorted(seen_levels['weak']))
print('в «Списке НПС», но не найдены в листах:', list(extra))
print('сайд-кики без статуса из второго файла:', missing_extra)
