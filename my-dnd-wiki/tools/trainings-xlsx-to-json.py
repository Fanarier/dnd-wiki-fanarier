# Тренировки и обучение из таблицы «Персонажи GATE H&A.xlsx» → server/seed/trainings.json.
# Сервер при запуске один раз вписывает их в листы (по id листа, а для тех, у кого листа пока нет, — по имени,
# когда такой лист появится). Цвет строки в Excel — статус: зелёный завершено, жёлтый в процессе, синий запланировано.
# Запуск: python tools/trainings-xlsx-to-json.py "путь/к/Персонажи GATE H&A.xlsx"
import json
import re
import secrets
import sys
from pathlib import Path

import openpyxl

SRC = sys.argv[1]
OUT = Path(__file__).resolve().parent.parent / 'server' / 'seed' / 'trainings.json'

# блок Excel (лист@строка заголовка) → лист на сайте. Вторые блоки на листах Хироши и Цуруи — Касуми и Долор
BY_ID = {
    'Хироши Мунэакира@37': 'sh719a8334f0',
    'Хироши Мунэакира@102': 'sh5eca616e79',  # Касуми Аои
    'Элиас Виан@37': 'sh427303dc72',
    'Цуруя Варабихимэ@42': 'shd6ba99bdaa',
    'Цуруя Варабихимэ@126': 'sha3810f2fe7',  # Долор
    'Лева Рутенхерц@42': 'sh1b786386f4',  # Löwe Rotesherz
    'Энди@38': 'shef6dc844d0',
}
FILL = {'FF93C47D': 'done', 'FFFFE599': 'active', 'FFFFD966': 'active', 'FF6FA8DC': 'planned'}


def note(c):
    if not c.comment:
        return ''
    t = c.comment.text.strip()
    if t.startswith('[Threaded comment]'):
        m = re.search(r'Comment:\s*(.*)', t, re.S)
        t = m.group(1).strip() if m else ''
    return re.sub(r'\s+', ' ', t)


def progress(v):
    s = str(v or '').strip()
    parts = re.split(r'\s*[/\\]\s*', s)
    if len(parts) == 2 and all(p.isdigit() for p in parts):
        return int(parts[0]), int(parts[1])
    # «5050» — пропущенная черта: половинки одинаковые
    if s.isdigit() and len(s) % 2 == 0 and s[:len(s) // 2] == s[len(s) // 2:]:
        return int(s[:len(s) // 2]), int(s[len(s) // 2:])
    return (int(s), 0) if s.isdigit() else (0, 0)


def stage(v):
    if v is None:
        return ''
    if isinstance(v, (int, float)):
        return str(int(v))
    return str(v).strip()


def rows_of(ws, r0):
    out = []
    r = r0 + 1
    while ws.cell(r, 1).value is not None or ws.cell(r, 2).value is not None:
        a, b, n = ws.cell(r, 1), ws.cell(r, 2), ws.cell(r, 3)
        cur, mx = progress(b.value)
        fill = str(a.fill.fgColor.rgb) if a.fill and a.fill.fill_type else ''
        want = FILL.get(fill, '')
        auto = 'done' if mx > 0 and cur >= mx else 'active'
        out.append({
            'Id': 't' + secrets.token_hex(5),
            'Name': str(a.value or '').strip() or 'Тренировка',
            'Note': note(a),
            'Cur': cur, 'Max': mx,
            'Stage': stage(n.value),
            'ProgNote': note(b),
            # статус пишем, только если он отличается от того, что посчитается сам
            'Status': want if want and want != auto else '',
        })
        r += 1
    return out


wb = openpyxl.load_workbook(SRC)
by_id, by_name = {}, {}
for ws in wb.worksheets:
    for (c0,) in ws.iter_rows(max_col=1):
        if isinstance(c0.value, str) and c0.value.strip().startswith('Тренировки и обучени'):
            key = f'{ws.title}@{c0.row}'
            rows = rows_of(ws, c0.row)
            if key in BY_ID:
                by_id[BY_ID[key]] = rows
            elif not any(k.startswith(ws.title + '@') for k in BY_ID):
                by_name[ws.title] = rows
            print(key, '→', BY_ID.get(key, ws.title), len(rows))

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps({'byId': by_id, 'byName': by_name}, ensure_ascii=False, indent=1), encoding='utf8')
print('записано', OUT)
