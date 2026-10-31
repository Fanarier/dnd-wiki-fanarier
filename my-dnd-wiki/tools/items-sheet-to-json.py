# Таблица уникальных предметов (Google Sheets «Предметы», вкладки по героям + «Общее») → server/seed/items.json
# и картинки предметов → server/seed/items-img/.
# Тексты берутся из CSV вкладок (как их видно в таблице: «35%», а не 0.35), оформление — из выгрузки xlsx:
# картинки (по строке, к какому предмету прицеплены), комментарии к ячейкам (→ примечание пункта),
# заливка ячейки (зелёная / жёлтая / красная / синяя → отметка пункта) и красный зачёркнутый текст (→ «утрачено»).
#   CSV:  https://docs.google.com/spreadsheets/d/<id>/export?format=csv&gid=<gid>  → <папка>/t<gid>.csv
#   xlsx: https://docs.google.com/spreadsheets/d/<id>/export?format=xlsx           → <папка>/items.xlsx
# Запуск: python tools/items-sheet-to-json.py <папка>
# Сервер вписывает каждую вещь один раз: в лист хозяина (если вещь уже там лежит — под тем же или похожим
# названием — ей добавляется карточка, иначе кладётся новая), «Общее» — в вики «Предметы».
import csv, hashlib, io, json, os, sys
import openpyxl
from PIL import Image

SRC = sys.argv[1]
ROOT = os.path.join(os.path.dirname(__file__), '..', 'server', 'seed')
OUT = os.path.join(ROOT, 'items.json')
IMG_DIR = os.path.join(ROOT, 'items-img')

# вкладка → хозяин (имя листа); на двоих — по каждой вещи отдельно (кто её носит в листе или по требованиям).
# Порядок — как вкладки в таблице (по нему же берутся листы xlsx)
TABS = {
    '0': None,  # Хироши и Касуми — см. BY_ITEM
    '1177847413': 'Элиас Виан',
    '587810015': None,  # Цуруя и Долор
    '1510971390': 'Löwe Rotesherz',
    '971934017': 'Белатор Светозаров',
    '443295819': 'Энди',
    '926668895': 'Азиэль Астарот',  # Азиэль и Азазэль: у Азазэля листа нет, вещи лежат у Азиэля
    '131642554': 'Орвилл Солес',
    '405056523': 'wiki',  # Общее
}
BY_ITEM = {
    'Катэкири Ясутсура': 'Хироши Мунэакира',
    'Перстень благочестия': 'Хироши Мунэакира',
    'Том астроманта': 'Касуми Аои',
    'Черный похититель звезд': 'Касуми Аои',
    'Последний вздох': 'Цуруя Варабихимэ',
    'Пояс драконьего гладиатора': 'Долор',
    'Кольцо здоровья': 'Долор',
}
# в листе вещь записана под другим названием
ALIAS = {
    'Катэкири Ясутсура': 'Оникири Ясутсура',
    'Пронзатель небес': 'Драконий лук',
    'Когти Алисфизы Астарот': 'Когти Астарот',
}
# последняя строка характеристик — чья работа (народ или мастер)
ORIGINS = {'Кицунэ', 'Драконы', 'Дворфы', 'Они', 'Древние', 'Тифлинги', 'Вампиры', 'Белатор'}
NONE = {'нет', ''}
# заливка ячейки → отметка: цвета темы Google (accent1..4) и те же цвета кодом
THEME_MARK = {4: 'blue', 5: 'red', 6: 'yellow', 7: 'green'}
RGB_MARK = {'4285F4': 'blue', '4A86E8': 'blue', 'EA4335': 'red', 'FF0000': 'red', 'CC0000': 'red', 'FBBC04': 'yellow', 'FFFF00': 'yellow', 'FF9900': 'yellow',
            '34A853': 'green', '00FF00': 'green', '6AA84F': 'green', '93C47D': 'green', 'B6D7A8': 'green'}
IMG_MAX = 1200


def cell(r, i):
    return r[i].strip() if i < len(r) else ''


def mark_of(c):
    """отметка пункта по оформлению ячейки xlsx"""
    f = c.font
    if f and f.strike:
        return 'gone'  # зачёркнуто (в таблице — красным): эффект утрачен
    fl = c.fill
    if fl is not None and fl.fill_type == 'solid':
        col = fl.fgColor
        if col.type == 'theme':
            return THEME_MARK.get(col.theme, '')
        if col.type == 'rgb' and col.rgb:
            return RGB_MARK.get(col.rgb[-6:].upper(), '')
    return ''


def note_of(c):
    return c.comment.text.strip() if c.comment and c.comment.text else ''


def save_img(data):
    im = Image.open(io.BytesIO(data))
    im.load()
    if max(im.size) > IMG_MAX:
        im.thumbnail((IMG_MAX, IMG_MAX), Image.LANCZOS)
    im = im.convert('RGBA') if im.mode in ('P', 'LA') else im
    buf = io.BytesIO()
    im.save(buf, 'WEBP', quality=86, method=6)
    raw = buf.getvalue()
    name = hashlib.sha1(raw).hexdigest()[:24] + '.webp'
    os.makedirs(IMG_DIR, exist_ok=True)
    with open(os.path.join(IMG_DIR, name), 'wb') as fh:
        fh.write(raw)
    return name


wb = openpyxl.load_workbook(os.path.join(SRC, 'items.xlsx'))
if len(wb.worksheets) != len(TABS):
    sys.exit(f'вкладок в xlsx {len(wb.worksheets)}, ждали {len(TABS)}')
for f in os.listdir(IMG_DIR) if os.path.isdir(IMG_DIR) else []:
    os.remove(os.path.join(IMG_DIR, f))

out = []
lost = []  # примечания, которые некуда положить
for (gid, owner), ws in zip(TABS.items(), wb.worksheets):
    rows = list(csv.reader(open(os.path.join(SRC, f't{gid}.csv'), encoding='utf8')))
    X = lambda r, c: ws.cell(row=r + 1, column=c + 1)  # ячейка xlsx по номерам CSV (с нуля)
    starts = [i for i, r in enumerate(rows) if cell(r, 6) == 'Название']
    imgs = sorted(((im.anchor._from.row, im) for im in ws._images), key=lambda x: x[0])
    for k, s in enumerate(starts):
        end = starts[k + 1] if k + 1 < len(starts) else max(len(rows), ws.max_row)
        block = rows[s:end]
        name = cell(block[0], 7)
        if not name:
            continue  # пустая заготовка
        if str(X(s, 7).value or '').strip() != name:
            sys.exit(f'{ws.title}: CSV и xlsx разошлись в строке {s + 1}')
        meta = {cell(r, 6): cell(r, 7) for r in block[1:7]}
        card = {'Type1': meta.get('Тип основной', ''), 'Type2': meta.get('Тип дополнительный', ''), 'Attune': meta.get('Сложность настройки', ''),
                'Kind': meta.get('Вид', ''), 'Origin': '', 'Note': note_of(X(s, 7)) or note_of(X(s, 6)), 'Chars': [], 'Comps': [], 'Reqs': [], 'Passive': [], 'Active': []}
        if card['Attune'].lower() == 'нет':
            card['Attune'] = ''
        for j in range(1, 7):
            for c in (6, 7):
                if note_of(X(s + j, c)):
                    lost.append(f'{name}: {cell(block[j], c)} — {note_of(X(s + j, c))[:60]}')
        for j, r in enumerate(block[1:], 1):
            for col, key in ((9, 'Comps'), (10, 'Reqs'), (12, 'Passive'), (13, 'Active')):
                v = cell(r, col)
                c = X(s + j, col)
                if v.lower() not in NONE:
                    e = {'T': v}
                    if note_of(c): e['N'] = note_of(c)
                    if mark_of(c): e['M'] = mark_of(c)
                    card[key].append(e)
                elif note_of(c):
                    lost.append(f'{name}: [{key}] «{v}» — {note_of(c)[:60]}')
        ch = next((i for i, r in enumerate(block) if cell(r, 6) == 'Характеристики'), None)
        if ch is not None:
            chars = []
            for j in range(ch + 1, len(block)):
                kk, vv = cell(block[j], 6), cell(block[j], 7)
                if not kk:
                    if vv or note_of(X(s + j, 7)):
                        lost.append(f'{name}: значение без названия «{vv}»')
                    continue
                notes = [n for n in (note_of(X(s + j, 6)), note_of(X(s + j, 7))) if n]
                e = {'K': kk, 'V': vv}
                if notes: e['N'] = '\n'.join(notes)
                m = mark_of(X(s + j, 6)) or mark_of(X(s + j, 7))
                if m: e['M'] = m
                chars.append(e)
            if chars and not chars[-1]['V'] and chars[-1]['K'] in ORIGINS and 'N' not in chars[-1]:
                card['Origin'] = chars.pop()['K']
            card['Chars'] = chars
        pic = [im for row, im in imgs if s <= row < end]
        img = save_img(pic[0]._data()) if pic else ''
        to = BY_ITEM.get(name, owner)
        if not to:
            sys.exit(f'нет хозяина: {name}')
        out.append({'key': f'{gid}:{name}', 'to': None if to == 'wiki' else to, 'shared': to == 'wiki', 'match': ALIAS.get(name, name),
                    'img': img, 'item': {'Name': name, 'Rarity': meta.get('Редкость', ''), 'Card': card}})

# одна и та же картинка у многих вещей — заглушка («?»), не переносим
uses = {}
for e in out:
    uses[e['img']] = uses.get(e['img'], 0) + 1
for e in out:
    if e['img'] and uses[e['img']] > 2:
        e['img'] = ''
for name, n in uses.items():
    if name and n > 2:
        os.remove(os.path.join(IMG_DIR, name))

json.dump(out, open(OUT, 'w', encoding='utf8'), ensure_ascii=False, indent=1)
print(len(out), 'предметов →', os.path.normpath(OUT))
for e in out:
    c = e['item']['Card']
    notes = bool(c['Note']) + sum('N' in x for k in ('Chars', 'Comps', 'Reqs', 'Passive', 'Active') for x in c[k])
    marks = sum('M' in x for k in ('Chars', 'Comps', 'Reqs', 'Passive', 'Active') for x in c[k])
    print(f"  {e['to'] or 'ВИКИ (общее)':22} {e['item']['Name']:28} {'🖼' if e['img'] else '  '} прим. {notes:2} отм. {marks}" + (f"  ⇐ «{e['match']}»" if e['match'] != e['item']['Name'] else ''))
for x in lost:
    print('  ! некуда:', x)
