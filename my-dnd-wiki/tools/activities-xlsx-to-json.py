# Журналы деятельностей героев из Google Таблицы (одна вкладка — один герой) → server/seed/activities.json.
# Перенос один раз; дальше игроки пишут деятельности на сайте («Деятельности» → «Журнал»).
#   python tools/activities-xlsx-to-json.py <файл .xlsx>
# Колонки: Деятельность | Часы | Выносливость | Тип деятельности | Статус | Результат | Дата.
# Строки-суммы (=SUM) и строки с одной датой пропускаем — суммы сайт считает сам.
import datetime
import hashlib
import json
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / 'server' / 'seed' / 'activities.json'
# вкладка → карточка героя
HEROES = {'Энди': 'hs01', 'Белатор': 'hs02', 'Джилл': 'hs03', 'Азиэль': 'hs04', 'Азазель': 'hs05', 'Лева': 'hs06',
          'Цуруя': 'hs07', 'Долор': 'hs08', 'Элиас': 'hs10', 'Хироши': 'hs11', 'Касуми': 'hs12'}
TYPES = {'Повторяемая': 'repeat', 'Уникальная': 'unique'}
STATUS = {'Учтено': 'done', 'Провал': 'fail', 'Условие': 'cond', 'Отменено': 'cancel'}


def text(v):
    if v is None:
        return ''
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(round(v, 3)).replace('.', ',')
    return str(v).strip()


wb = openpyxl.load_workbook(sys.argv[1])
vals = openpyxl.load_workbook(sys.argv[1], data_only=True)  # посчитанные значения формул («=SUM(89+82…)» → 548)
logs, total = {}, 0
for ws in wb.worksheets:
    wv = vals[ws.title]
    hero = HEROES.get(ws.title.strip())
    if not hero:
        print('нет героя для вкладки', ws.title)
        continue
    rows, date = [], ''
    for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
        a, b, c, d, e, f, g = [x.value for x in r[:7]]
        if isinstance(b, str) and b.startswith('='):
            continue  # строка-сумма
        # формула в ячейке — берём её значение
        a, b, c, d, e, f, g = [wv.cell(i, k + 1).value if isinstance(v, str) and v.startswith('=') else v for k, v in enumerate((a, b, c, d, e, f, g))]
        if isinstance(g, datetime.datetime):
            date = g.date().isoformat()  # строка без даты — та же сессия, что выше
        name = text(a)
        if not name or (name == '???' and not any((b, c, d, e, f))):
            continue
        rows.append({
            'id': 'a' + hashlib.md5(f'{hero}|{i}|{name}'.encode('utf8')).hexdigest()[:10],
            'date': date, 'name': name[:600], 'hours': text(b)[:30], 'stamina': text(c)[:30],
            'type': TYPES.get(text(d), ''), 'status': STATUS.get(text(e), 'pending'), 'result': text(f)[:2000],
        })
    logs[hero] = rows
    total += len(rows)
    print(f'{ws.title}: {len(rows)}')
OUT.write_text(json.dumps({'logs': logs}, ensure_ascii=False), encoding='utf8')
print('всего', total)
