# Квесты и хроника сюжета из таблицы «Анкария GATE H&A.xlsx» (листы «Квесты» и «Сюжет») →
# server/seed/campaign.json. Сайт берёт этот файл один раз, дальше всё правит мастер на странице «Квесты».
#   python tools/campaign-xlsx-to-json.py "путь/к/таблице.xlsx"
import json
import sys
import datetime
from pathlib import Path

import openpyxl

src = sys.argv[1]
wb = openpyxl.load_workbook(src, data_only=True)


def txt(v):
    if v is None:
        return ''
    if isinstance(v, datetime.datetime):
        return v.strftime('%d.%m.%Y')
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


FIX = {'Старна Зайцегонов': 'Страна Зайцегонов'}

# ---- квесты ----
ws = wb['Квесты']
quests = []
for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True)):
    if not r[0]:
        continue
    name, diff, kind, region, status, taker, loc, desc, source = (list(r) + [None] * 9)[:9]
    quests.append({
        'id': f'cq{i:03d}', 'name': txt(name), 'difficulty': txt(diff), 'type': txt(kind),
        'region': FIX.get(txt(region), txt(region)), 'status': txt(status), 'taker': txt(taker),
        'location': txt(loc), 'desc': txt(desc), 'source': txt(source), 'hidden': False
    })

# ---- сюжет: Том (строка 1) → Акт (объединённые строки) → главы ----
ws = wb['Сюжет']
volumes = [txt(ws['A1'].value)] + ([txt(ws['G1'].value)] if ws['G1'].value else [])
chapters = []
act = ''
for row in ws.iter_rows(min_row=3, max_row=80, max_col=5, values_only=True):
    a, status, result, dworld, dreal = row
    if not a:
        continue
    a = txt(a)
    if a.startswith('Акт') and not status:
        act = a
        continue
    chapters.append({
        'id': f'ch{len(chapters):03d}', 'volume': volumes[0], 'act': act, 'title': a, 'status': txt(status),
        'result': txt(result), 'dateWorld': txt(dworld), 'dateReal': txt(dreal), 'hidden': False
    })

out = Path(__file__).resolve().parent.parent / 'server' / 'seed' / 'campaign.json'
out.write_text(json.dumps({'volumes': volumes, 'quests': quests, 'chapters': chapters}, ensure_ascii=False, indent=1), encoding='utf-8')
print(f'квестов: {len(quests)}, глав: {len(chapters)}, томов: {len(volumes)} → {out}')
