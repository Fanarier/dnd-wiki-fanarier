"""Бестиарий из таблицы «Бестиарий H&A.xlsx» → server/seed/bestiary.json и картинки в server/seed/bestiary-img/.

Каждый лист — группа существ («Големы»). Моб — блок от строки с «Имя» в колонке F: слева пары «поле — значение»
(F→G: размер, тип, хиты, защита…; H→I: характеристики, зрение, слух, аура), справа списки по колонкам (K владения,
L спасброски, M сопротивления, N уязвимости, O иммунитеты, Q пассивные, R атаки, S активные, T легендарные, U мифические),
картинка привязана к колонке B в той же строке.

Запуск: python tools/bestiary-xlsx-to-json.py <файл.xlsx> [папка вывода, по умолчанию server/seed]
"""
import json, os, re, sys
import openpyxl

src = sys.argv[1]
out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(__file__), '..', 'server', 'seed')
img_dir = os.path.join(out_dir, 'bestiary-img')
os.makedirs(img_dir, exist_ok=True)

LISTS = {'K': 'prof', 'L': 'saves', 'M': 'resist', 'N': 'vuln', 'O': 'immune', 'Q': 'passive', 'R': 'attacks', 'S': 'active', 'T': 'legendary', 'U': 'mythic'}
STATS = ['Сила', 'Телосложение', 'Интеллект', 'Харизма', 'Ловкость', 'Мудрость']
SENSES = ['Зрение', 'Слух', 'Аура']
SPEEDS = ['Базовая скорость', 'Полет', 'Плавание']


def val(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return str(v).strip() if isinstance(v, str) else v


wb = openpyxl.load_workbook(src, data_only=True)
mobs = []
# одна и та же картинка у нескольких мобов — заглушка («?»), её не переносим
import hashlib
seen_hash, img_data = {}, {}  # картинку из файла можно прочитать только один раз — запоминаем
for ws in wb.worksheets:
    for im in getattr(ws, '_images', []):
        img_data[id(im)] = im._data()
        h = hashlib.md5(img_data[id(im)]).hexdigest()
        seen_hash[h] = seen_hash.get(h, 0) + 1
for ws in wb.worksheets:
    images = {}
    for im in getattr(ws, '_images', []):
        images[im.anchor._from.row + 1] = im
    starts = [r for r in range(1, ws.max_row + 1) if ws[f'F{r}'].value == 'Имя']
    for i, r in enumerate(starts):
        end = starts[i + 1] if i + 1 < len(starts) else ws.max_row + 1
        left, mid = {}, {}
        for rr in range(r, min(end, r + 25)):
            k, v = ws[f'F{rr}'].value, ws[f'G{rr}'].value
            if k: left[str(k).strip()] = val(v)
            k, v = ws[f'H{rr}'].value, ws[f'I{rr}'].value
            if k: mid[str(k).strip()] = val(v)
        lists = {}
        for col, key in LISTS.items():
            items = []
            for rr in range(r + 1, min(end, r + 25)):
                v = ws[f'{col}{rr}'].value
                if v is None or str(v).strip() == '':
                    break
                items.append(str(v).strip())
            lists[key] = items
        name = str(left.get('Имя', '')).strip()
        mob = {
            'name': name, 'group': ws.title,
            'size': left.get('Размер', ''), 'type': left.get('Тип', ''), 'race': left.get('Раса', ''), 'align': left.get('Мировоззрение', ''),
            'stats': [{'k': k, 'v': mid.get(k, '')} for k in STATS],
            'hp': left.get('Хиты', ''), 'def': left.get('Защита', ''), 'eva': left.get('Уклонение', ''),
            'senses': [{'k': k, 'v': mid.get(k, '')} for k in SENSES],
            'speed': [{'k': k, 'v': left.get(k, '')} for k in SPEEDS],
            'bm': left.get('БМ', ''), 'langs': left.get('Языки', ''), 'exp': left.get('Опыт', ''), 'danger': left.get('Опасность', ''),
            **lists, 'desc': '', 'img': ''
        }
        im = images.get(r)
        if im is not None and seen_hash.get(hashlib.md5(img_data[id(im)]).hexdigest(), 0) > 1:
            im = None
        if im is not None:
            data = img_data[id(im)]
            ext = 'png' if data[:4] == b'\x89PNG' else 'jpg' if data[:2] == b'\xff\xd8' else 'webp' if data[8:12] == b'WEBP' else 'png'
            slug = re.sub(r'[^a-z0-9]+', '-', ws.title.encode('ascii', 'ignore').decode().lower()).strip('-') or 'mob'
            fn = f'mob-{len(mobs) + 1:03d}-{slug}.{ext}'
            with open(os.path.join(img_dir, fn), 'wb') as f:
                f.write(data)
            mob['img'] = fn
        mobs.append(mob)

with open(os.path.join(out_dir, 'bestiary.json'), 'w', encoding='utf-8') as f:
    json.dump({'mobs': mobs}, f, ensure_ascii=False, indent=1)
print(f'мобов: {len(mobs)}, с картинкой: {sum(1 for m in mobs if m["img"])}')
for m in mobs:
    print(' ', m['group'], '·', m['name'], '·', m['danger'], '·', m['img'])
